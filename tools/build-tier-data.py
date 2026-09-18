#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel -> <slug>-data.js  (Гишүүн / Бизнес / Мастер хуудсууд)
=============================================================
Гишүүнчлэлийн системээс татсан Excel-ээс (14 багана, sheet бүр нэг
ангилал) эрсдэлтэй гишүүдийн жагсаалтыг уншиж, хуудасны өгөгдлийн
файлыг үүсгэнэ. Хэд хэдэн sheet-ийг нэг хуудсанд нэгтгэж болно.

Хэрэглээ:
    py tools/build-tier-data.py <excel> <slug> <sheet> [<sheet> ...]

Жишээ:
    py tools/build-tier-data.py "гишүүд.xlsx" rtqfy9q7pd Member
    py tools/build-tier-data.py "гишүүд.xlsx" sb8nhuuzep "Full member" Start
    py tools/build-tier-data.py "гишүүд.xlsx" ywufhhbsci Master

Багануудыг НЭРЭЭР нь олдог (1-р мөр толгой). Хэрэглэдэг баганууд:
    membership number · name · registration date · Name Position
    Sponsor number · Sponsor's name · Referral number · recommended person
    Уригчийн уригч ID · Уригчийн уригч нэр

ХУВИЙН МЭДЭЭЛЛИЙН ХАМГААЛАЛТ: "cell phone", "Resident registration number"
зэрэг багануудыг ЗОРИУДААР уншдаггүй — вэб рүү хэзээ ч гарахгүй.

Дараа нь:
    git add <slug>-data.js
    git commit -m "Зарлал шинэчлэв"
    git push origin main
"""

import collections
import datetime
import io
import json
import os
import re
import sys

import openpyxl

MISSING = ("", "#N/A", "#N/A!", "#VALUE!", "#REF!", "NONE", "NULL")

# гаралтын талбар -> Excel-ийн толгойн нэр (жижиг үсгээр)
COLUMNS = {
    "id":        "membership number",
    "name":      "name",
    "date":      "registration date",
    "position":  "name position",
    "spon_id":   "sponsor number",
    "spon_name": "sponsor's name",
    "inv_id":    "referral number",
    "inv_name":  "recommended person",
    "up_id":     "уригчийн уригч id",
    "up_name":   "уригчийн уригч нэр",
}

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def clean(v):
    return re.sub(r"\s+", " ", str(v).strip()) if v is not None else ""


def missing(v):
    return v.upper() in MISSING


def header_map(ws):
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {clean(h).lower(): i for i, h in enumerate(first) if h is not None}
    col = {}
    for field, want in COLUMNS.items():
        if want not in headers:
            raise SystemExit('ERROR: sheet "%s" дээр "%s" багана олдсонгүй. Толгой: %s'
                             % (ws.title, want, list(headers)))
        col[field] = headers[want]
    return col


def main():
    if len(sys.argv) < 4:
        print(__doc__)
        return 1
    src, slug, sheets = sys.argv[1], sys.argv[2], sys.argv[3:]
    out = os.path.join(HERE, slug + "-data.js")

    wb = openpyxl.load_workbook(src, data_only=True)
    people, index, rows = [], {}, []
    per_sheet = collections.OrderedDict()

    def ref(pid, pname):
        if missing(pid):
            return -1
        if pid not in index:
            index[pid] = len(people)
            people.append([pid, pname])
        elif not people[index[pid]][1] and pname:
            people[index[pid]][1] = pname
        return index[pid]

    for sheet in sheets:
        if sheet not in wb.sheetnames:
            raise SystemExit('ERROR: "%s" нэртэй sheet алга. Байгаа нь: %s' % (sheet, wb.sheetnames))
        ws = wb[sheet]
        col = header_map(ws)
        n = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or col["id"] >= len(r) or not r[col["id"]]:
                continue
            g = lambda f: clean(r[col[f]]) if col[f] < len(r) else ""
            rows.append([
                g("id"), g("name"), g("date"),
                ref(g("inv_id"), g("inv_name")),
                ref(g("spon_id"), g("spon_name")),
                ref(g("up_id"), g("up_name")),
                g("position"),
            ])
            n += 1
        per_sheet[sheet] = n

    dupes = len(rows) - len(set(r[0] for r in rows))
    if dupes:
        print("АНХААР: %d давхардсан гишүүний ID байна." % dupes)

    data = {
        "updated": datetime.date.today().isoformat(),
        "total": len(rows),
        "people": people,
        "rows": rows,
    }
    js = ("/* SKINDOX - Зарлах жагсаалт (auto-generated, do not edit by hand) */\n"
          "window.SKINDOX_ZARLAL=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n")
    with io.open(out, "w", encoding="utf-8") as f:
        f.write(js)

    print("OK -> %s (%d bytes)" % (out, os.path.getsize(out)))
    print("   sheet бүрээр:", dict(per_sheet))
    print("   гишүүн: %d | нэрсийн сан: %d" % (len(rows), len(people)))
    print("   ангилал:", dict(collections.Counter(r[6] for r in rows)))
    print("   уригчтай: %d | спонсортой: %d | уригчийн уригчтай: %d"
          % tuple(sum(1 for r in rows if r[i] >= 0) for i in (3, 4, 5)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
